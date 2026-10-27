"""Exporters: formatted Excel workbook, CSV and JSON."""

from __future__ import annotations

import csv
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, Optional

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .models import COLUMNS, Book

HEADERS = {
    "title": "Title",
    "category": "Category",
    "price": "Price (GBP)",
    "rating": "Rating (1-5)",
    "in_stock": "In stock",
    "stock_count": "Stock count",
    "upc": "UPC",
    "num_reviews": "Reviews",
    "url": "Product URL",
    "image_url": "Image URL",
    "description": "Description",
}
WIDTHS = {
    "title": 48,
    "category": 20,
    "price": 14,
    "rating": 15,
    "in_stock": 12,
    "stock_count": 15,
    "upc": 18,
    "num_reviews": 12,
    "url": 40,
    "image_url": 40,
    "description": 60,
}
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF")
GBP = '"£"#,##0.00'
CHART_TOP_N = 15


def _columns(books: list[Book]) -> list[str]:
    """Drop detail-only columns when no book has them (keeps sheets tidy)."""
    detail = ["stock_count", "upc", "num_reviews", "description"]
    has_details = any(b.upc is not None for b in books)
    return [c for c in COLUMNS if has_details or c not in detail]


def _style_header(ws, ncols: int, row: int = 1) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")


def write_excel(
    books: list[Book],
    path: Path,
    source_url: str = "",
    filters_desc: str = "",
    scraped_at: Optional[datetime] = None,
) -> Path:
    """Write a client-ready workbook.

    Sheets:
      * Books    - one row per book, as an Excel table (filter/sort buttons),
                   frozen header, GBP number format, clickable links.
      * Summary  - per-category count / average / min / max price and average
                   rating, written as live formulas over the Books table,
                   sorted by size, plus a bar chart of the top categories.
      * About    - when and how the data was collected.
    """
    scraped_at = scraped_at or datetime.now(timezone.utc)
    cols = _columns(books)
    wb = Workbook()

    # ---------------------------------------------------------------- Books
    ws = wb.active
    ws.title = "Books"
    ws.append([HEADERS[c] for c in cols])
    for b in books:
        d = b.to_dict()
        row = []
        for c in cols:
            v = d[c]
            if c == "in_stock":
                v = "Yes" if v else "No"
            row.append(v)
        ws.append(row)

    n = len(books)
    last_col = get_column_letter(len(cols))
    for i, c in enumerate(cols, 1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = WIDTHS[c]
        if c == "price":
            for r in range(2, n + 2):
                ws[f"{letter}{r}"].number_format = GBP
        if c in ("url", "image_url"):
            for r in range(2, n + 2):
                cell = ws[f"{letter}{r}"]
                if cell.value:
                    cell.hyperlink = cell.value
                    cell.font = Font(color="0563C1", underline="single")
        if c == "description":
            for r in range(2, n + 2):
                ws[f"{letter}{r}"].alignment = Alignment(wrap_text=False)
    ws.freeze_panes = "A2"
    if n:
        table = Table(displayName="Books", ref=f"A1:{last_col}{n + 1}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    # Style the header explicitly too: LibreOffice and Google Sheets ignore
    # Excel table styles, and this keeps the header readable everywhere.
    _style_header(ws, len(cols))

    # -------------------------------------------------------------- Summary
    sm = wb.create_sheet("Summary")
    sm.append(["Category", "Books", "Avg price", "Min price", "Max price", "Avg rating"])
    _style_header(sm, 6)
    cat_col = get_column_letter(cols.index("category") + 1)
    price_col = get_column_letter(cols.index("price") + 1)
    rate_col = get_column_letter(cols.index("rating") + 1)
    rng = lambda col: f"Books!${col}$2:${col}${max(n + 1, 2)}"  # noqa: E731
    # Biggest categories first; ties alphabetical.
    counts = Counter(b.category for b in books)
    categories = sorted(counts, key=lambda c: (-counts[c], c))
    for i, cat in enumerate(categories, start=2):
        sm.cell(row=i, column=1, value=cat)
        sm.cell(row=i, column=2, value=f"=COUNTIF({rng(cat_col)},A{i})")
        sm.cell(row=i, column=3, value=f"=AVERAGEIF({rng(cat_col)},A{i},{rng(price_col)})")
        sm.cell(row=i, column=4, value=f"=_xlfn.MINIFS({rng(price_col)},{rng(cat_col)},A{i})")
        sm.cell(row=i, column=5, value=f"=_xlfn.MAXIFS({rng(price_col)},{rng(cat_col)},A{i})")
        sm.cell(row=i, column=6, value=f"=AVERAGEIF({rng(cat_col)},A{i},{rng(rate_col)})")
        for col in (3, 4, 5):
            sm.cell(row=i, column=col).number_format = GBP
        sm.cell(row=i, column=6).number_format = "0.00"
    total_row = len(categories) + 2
    sm.cell(row=total_row, column=1, value="All").font = Font(bold=True)
    sm.cell(row=total_row, column=2, value=f"=SUM(B2:B{total_row - 1})").font = Font(bold=True)
    sm.cell(row=total_row, column=3, value=f"=AVERAGE({rng(price_col)})").number_format = GBP
    sm.cell(row=total_row, column=6, value=f"=AVERAGE({rng(rate_col)})").number_format = "0.00"
    for letter, width in zip("ABCDEF", (24, 10, 13, 13, 13, 13)):
        sm.column_dimensions[letter].width = width
    sm.freeze_panes = "A2"

    if categories:
        chart = BarChart()
        top = min(CHART_TOP_N, len(categories))
        chart.type = "bar"
        chart.title = f"Top {top} categories by number of books"
        chart.y_axis.title = "Books"
        chart.legend = None
        chart.height = max(7.5, 0.9 * top)
        chart.width = 18
        chart.x_axis.scaling.orientation = "maxMin"  # largest bar at the top
        chart.x_axis.delete = False  # newer openpyxl hides axes by default
        chart.y_axis.delete = False
        data = Reference(sm, min_col=2, min_row=1, max_row=top + 1)
        labels = Reference(sm, min_col=1, min_row=2, max_row=top + 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(labels)
        sm.add_chart(chart, "H2")

    # ---------------------------------------------------------------- About
    ab = wb.create_sheet("About")
    rows = [
        ("Source", source_url),
        ("Scraped at (UTC)", scraped_at.strftime("%Y-%m-%d %H:%M")),
        ("Books", n),
        ("Categories", len(categories)),
        ("Filters", filters_desc or "none"),
        ("Generated by", "bookscraper (https://github.com/ritanjay/web-to-excel-scraper)"),
    ]
    for k, v in rows:
        ab.append([k, v])
        ab.cell(row=ab.max_row, column=1).font = Font(bold=True)
    ab.column_dimensions["A"].width = 18
    ab.column_dimensions["B"].width = 70

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def write_csv(books: Iterable[Book], path: Path) -> Path:
    books = list(books)
    cols = _columns(books)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    # utf-8-sig so Excel on Windows shows "£" and accents correctly.
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        for b in books:
            w.writerow(b.to_dict())
    return path


def write_json(books: Iterable[Book], path: Path) -> Path:
    books = list(books)
    cols = _columns(books)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    data = [{c: b.to_dict()[c] for c in cols} for b in books]
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return path
