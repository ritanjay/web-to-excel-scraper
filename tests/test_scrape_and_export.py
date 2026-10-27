import csv
import json

import pytest
from openpyxl import load_workbook

from bookscraper.cli import main
from bookscraper.export import write_csv, write_excel, write_json
from bookscraper.scrape import Filters, scrape

from .conftest import BASE


def test_scrape_walks_categories_and_pagination(site):
    result = scrape(site)
    assert result.pages_visited == 3
    assert result.errors == []
    assert result.categories == ["Travel", "Poetry"]
    titles = [(b.category, b.title) for b in result.books]
    assert titles == sorted(titles)  # sorted by category, then title
    assert len(result.books) == 5
    assert {b.category for b in result.books} == {"Travel", "Poetry"}


def test_filters(site):
    result = scrape(site, filters=Filters(min_rating=3, max_price=50, in_stock_only=True))
    assert [b.title for b in result.books] == [
        "Shakespeare's Sonnets",
        "Full Moon over Noah’s Ark: An Odyssey to Mount Ararat and Beyond",
    ]


def test_category_selection_is_case_insensitive(site):
    result = scrape(site, filters=Filters(categories=["poetry"]))
    assert {b.category for b in result.books} == {"Poetry"}
    assert not any("travel" in u for u in site.requested)


def test_unknown_category_lists_available(site):
    with pytest.raises(ValueError, match="Available: Poetry, Travel"):
        scrape(site, filters=Filters(categories=["Cooking"]))


def test_max_pages_stops_early(site):
    result = scrape(site, max_pages=1)
    assert result.pages_visited == 1
    assert len(result.books) == 2


def test_details_fill_product_fields(site):
    result = scrape(site, filters=Filters(categories=["Poetry"]), details=True)
    book = result.books[0]
    assert book.upc == "a897fe39b1053632"
    assert book.stock_count == 22
    assert book.description.startswith("It's hard")


def test_broken_page_is_reported_not_fatal(site):
    site.pages.pop("https://books.toscrape.com/catalogue/category/books/travel_2/page-2.html")
    result = scrape(site)
    assert len(result.errors) == 1
    assert len(result.books) == 4  # everything else still collected


def test_excel_workbook_structure(site, tmp_path):
    books = scrape(site).books
    path = write_excel(books, tmp_path / "out.xlsx", source_url=BASE, filters_desc="none")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Books", "Summary", "About"]

    ws = wb["Books"]
    header = [c.value for c in ws[1]]
    assert header[:5] == ["Title", "Category", "Price (GBP)", "Rating (1-5)", "In stock"]
    assert "UPC" not in header  # detail columns hidden when not scraped
    assert ws.max_row == 6
    assert ws.freeze_panes == "A2"
    assert "Books" in ws.tables
    assert ws["C2"].number_format == '"£"#,##0.00'
    assert ws["F2"].hyperlink.target.startswith(BASE)

    sm = wb["Summary"]
    assert [sm.cell(row=r, column=1).value for r in range(2, 4)] == ["Poetry", "Travel"]
    assert sm["B2"].value.startswith("=COUNTIF(Books!$B$2:$B$6")
    assert sm["A4"].value == "All"
    assert len(sm._charts) == 1


def test_excel_formulas_compute_correctly(site, tmp_path):
    """Recalculate with LibreOffice (if installed) and check the summary numbers."""
    import shutil
    import subprocess

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        pytest.skip("LibreOffice not installed")
    books = scrape(site).books
    src = write_excel(books, tmp_path / "in.xlsx")
    outdir = tmp_path / "calc"
    subprocess.run(
        [soffice, "--headless", "--convert-to", "xlsx", "--outdir", str(outdir), str(src)],
        check=True,
        capture_output=True,
        timeout=120,
    )
    sm = load_workbook(outdir / "in.xlsx", data_only=True)["Summary"]
    rows = {
        sm.cell(row=r, column=1).value: [sm.cell(row=r, column=c).value for c in range(2, 7)]
        for r in range(2, 5)
    }
    assert rows["Poetry"][0] == 2
    assert rows["Poetry"][1] == pytest.approx((51.77 + 20.66) / 2)
    assert rows["Travel"][2] == pytest.approx(44.34)  # min
    assert rows["Travel"][3] == pytest.approx(49.43)  # max
    assert rows["All"][0] == 5


def test_details_columns_appear_when_scraped(site, tmp_path):
    books = scrape(site, filters=Filters(categories=["Poetry"]), details=True).books
    ws = load_workbook(write_excel(books, tmp_path / "d.xlsx"))["Books"]
    assert "UPC" in [c.value for c in ws[1]]


def test_csv_and_json(site, tmp_path):
    books = scrape(site).books
    c = write_csv(books, tmp_path / "b.csv")
    j = write_json(books, tmp_path / "b.json")
    raw = c.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")  # BOM so Excel opens UTF-8 correctly
    rows = list(csv.DictReader(c.open(encoding="utf-8-sig")))
    assert len(rows) == 5 and rows[0]["category"] == "Poetry"
    data = json.loads(j.read_text(encoding="utf-8"))
    assert data[0]["price"] == pytest.approx(20.66) or data[0]["price"] == pytest.approx(51.77)
    assert "upc" not in data[0]


def test_empty_result_still_writes_valid_workbook(tmp_path):
    wb = load_workbook(write_excel([], tmp_path / "e.xlsx"))
    assert wb["Books"].max_row == 1


def test_cli_end_to_end(site, tmp_path, capsys):
    out = tmp_path / "cli.xlsx"
    code = main(
        ["-o", str(out), "--csv", str(tmp_path / "cli.csv"), "--min-rating", "4"],
        fetcher=_counting(site),
    )
    assert code == 0
    assert out.exists() and (tmp_path / "cli.csv").exists()
    err = capsys.readouterr().err
    assert "Done: 2 books from 3 pages" in err
    about = load_workbook(out)["About"]
    assert about["B5"].value == "rating>=4"


def test_cli_unknown_category_exit_code(site, tmp_path, capsys):
    code = main(["-o", str(tmp_path / "x.xlsx"), "-c", "Nope"], fetcher=_counting(site))
    assert code == 2
    assert "unknown category" in capsys.readouterr().err


def _counting(site):
    site.requests_made = 0
    site.cache_hits = 0
    return site
